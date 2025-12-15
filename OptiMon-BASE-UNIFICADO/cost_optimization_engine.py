#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
🎯 OptiMon Cost Optimization Engine v3.1.0
======================================================
Motor de Recomendaciones Inteligentes para Optimización de Costos
Analiza patrones de uso y genera recomendaciones automáticas
"""

import json
import logging
import boto3
from datetime import datetime, timedelta
from typing import Dict, List, Any
from pathlib import Path
import requests
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    XLSX_AVAILABLE = True
except ImportError:
    XLSX_AVAILABLE = False

# Configurar logging
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

class CostOptimizationEngine:
    """Motor de optimización de costos para recursos en la nube"""
    
    def __init__(self):
        # Precios AWS EC2 On-Demand (USD/hora) - Región us-east-1
        # Fuente: AWS Pricing Calculator (Octubre 2025)
        self.aws_pricing = {
            # Familia T3 (Burstable Performance)
            't3.nano': {'price': 0.0052, 'vcpu': 2, 'memory': 0.5, 'family': 'General Purpose'},
            't3.micro': {'price': 0.0104, 'vcpu': 2, 'memory': 1, 'family': 'General Purpose'},
            't3.small': {'price': 0.0208, 'vcpu': 2, 'memory': 2, 'family': 'General Purpose'},
            't3.medium': {'price': 0.0416, 'vcpu': 2, 'memory': 4, 'family': 'General Purpose'},
            't3.large': {'price': 0.0832, 'vcpu': 2, 'memory': 8, 'family': 'General Purpose'},
            't3.xlarge': {'price': 0.1664, 'vcpu': 4, 'memory': 16, 'family': 'General Purpose'},
            # Familia M5 (General Purpose)
            'm5.large': {'price': 0.096, 'vcpu': 2, 'memory': 8, 'family': 'General Purpose'},
            'm5.xlarge': {'price': 0.192, 'vcpu': 4, 'memory': 16, 'family': 'General Purpose'},
            'm5.2xlarge': {'price': 0.384, 'vcpu': 8, 'memory': 32, 'family': 'General Purpose'},
            # Familia C5 (Compute Optimized)
            'c5.large': {'price': 0.085, 'vcpu': 2, 'memory': 4, 'family': 'Compute Optimized'},
            'c5.xlarge': {'price': 0.17, 'vcpu': 4, 'memory': 8, 'family': 'Compute Optimized'},
            # Familia R5 (Memory Optimized)
            'r5.large': {'price': 0.126, 'vcpu': 2, 'memory': 16, 'family': 'Memory Optimized'},
            'r5.xlarge': {'price': 0.252, 'vcpu': 4, 'memory': 32, 'family': 'Memory Optimized'}
        }
        
        # Precios Azure VMs (USD/hora) - Región East US
        # Fuente: Azure Pricing Calculator (Octubre 2025)
        self.azure_pricing = {
            # Familia B (Burstable)
            'Standard_B1ls': {'price': 0.0052, 'vcpu': 1, 'memory': 0.5, 'family': 'Burstable'},
            'Standard_B1s': {'price': 0.0104, 'vcpu': 1, 'memory': 1, 'family': 'Burstable'},
            'Standard_B1ms': {'price': 0.0208, 'vcpu': 1, 'memory': 2, 'family': 'Burstable'},
            'Standard_B2s': {'price': 0.0416, 'vcpu': 2, 'memory': 4, 'family': 'Burstable'},
            'Standard_B2ms': {'price': 0.0832, 'vcpu': 2, 'memory': 8, 'family': 'Burstable'},
            # Familia D (General Purpose)
            'Standard_D2s_v3': {'price': 0.096, 'vcpu': 2, 'memory': 8, 'family': 'General Purpose'},
            'Standard_D4s_v3': {'price': 0.192, 'vcpu': 4, 'memory': 16, 'family': 'General Purpose'},
            'Standard_D8s_v3': {'price': 0.384, 'vcpu': 8, 'memory': 32, 'family': 'General Purpose'},
            # Familia F (Compute Optimized)
            'Standard_F2s_v2': {'price': 0.085, 'vcpu': 2, 'memory': 4, 'family': 'Compute Optimized'},
            'Standard_F4s_v2': {'price': 0.17, 'vcpu': 4, 'memory': 8, 'family': 'Compute Optimized'}
        }
        
        # Información de fuentes de precios
        self.pricing_sources = {
            'aws': 'AWS Pricing Calculator - us-east-1 (Oct 2025)',
            'azure': 'Azure Pricing Calculator - East US (Oct 2025)',
            'last_updated': '2025-10-26'
        }
        
        self.recommendations_cache = []
        
    def analyze_resource_utilization(self, provider: str, instance_id: str, 
                                   metrics_data: Dict) -> Dict[str, Any]:
        """Analizar utilización de recursos de una instancia"""
        try:
            # Extraer métricas principales
            cpu_avg = metrics_data.get('cpu_average', 0)
            cpu_max = metrics_data.get('cpu_max', 0)
            memory_avg = metrics_data.get('memory_average', 0)
            memory_max = metrics_data.get('memory_max', 0)
            network_in = metrics_data.get('network_in', 0)
            network_out = metrics_data.get('network_out', 0)
            
            # Calcular puntuaciones de utilización
            cpu_score = self._calculate_utilization_score(cpu_avg, cpu_max)
            memory_score = self._calculate_utilization_score(memory_avg, memory_max)
            
            # Determinar patrón de uso
            usage_pattern = self._determine_usage_pattern(metrics_data)
            
            return {
                'instance_id': instance_id,
                'provider': provider,
                'cpu_utilization': {
                    'average': cpu_avg,
                    'max': cpu_max,
                    'score': cpu_score
                },
                'memory_utilization': {
                    'average': memory_avg,
                    'max': memory_max,
                    'score': memory_score
                },
                'network_usage': {
                    'inbound_gb': network_in / (1024**3),
                    'outbound_gb': network_out / (1024**3)
                },
                'usage_pattern': usage_pattern,
                'analysis_timestamp': datetime.now().isoformat()
            }
            
        except Exception as e:
            logger.error(f"Error analizando utilización {instance_id}: {e}")
            return {}
    
    def generate_cost_recommendations(self, analysis_data: Dict, 
                                    current_instance_type: str) -> List[Dict]:
        """Generar recomendaciones de optimización de costos"""
        recommendations = []
        
        try:
            provider = analysis_data.get('provider', 'aws')
            cpu_score = analysis_data.get('cpu_utilization', {}).get('score', 0)
            memory_score = analysis_data.get('memory_utilization', {}).get('score', 0)
            usage_pattern = analysis_data.get('usage_pattern', 'unknown')
            
            # Recomendación de rightsizing
            rightsizing_rec = self._generate_rightsizing_recommendation(
                provider, current_instance_type, cpu_score, memory_score
            )
            if rightsizing_rec:
                recommendations.append(rightsizing_rec)
            
            # Recomendación de scheduling
            scheduling_rec = self._generate_scheduling_recommendation(
                analysis_data, usage_pattern
            )
            if scheduling_rec:
                recommendations.append(scheduling_rec)
            
            # Recomendación de reserved instances
            reserved_rec = self._generate_reserved_instances_recommendation(
                provider, current_instance_type, usage_pattern
            )
            if reserved_rec:
                recommendations.append(reserved_rec)
            
            # Recomendación de spot instances
            spot_rec = self._generate_spot_instances_recommendation(
                provider, current_instance_type, usage_pattern
            )
            if spot_rec:
                recommendations.append(spot_rec)
                
        except Exception as e:
            logger.error(f"Error generando recomendaciones: {e}")
        
        return recommendations
    
    def _calculate_utilization_score(self, avg_usage: float, max_usage: float) -> str:
        """Calcular puntuación de utilización"""
        if avg_usage < 10:
            return "muy_bajo"
        elif avg_usage < 30:
            return "bajo"
        elif avg_usage < 60:
            return "optimo"
        elif avg_usage < 80:
            return "alto"
        else:
            return "muy_alto"
    
    def _determine_usage_pattern(self, metrics_data: Dict) -> str:
        """Determinar patrón de uso basado en métricas históricas"""
        # Análisis simplificado del patrón
        cpu_avg = metrics_data.get('cpu_average', 0)
        cpu_max = metrics_data.get('cpu_max', 0)
        
        variance = cpu_max - cpu_avg
        
        if cpu_avg < 10:
            return "idle"
        elif variance > 50:
            return "variable"
        elif cpu_avg > 70:
            return "intensivo"
        else:
            return "estable"
    
    def _generate_rightsizing_recommendation(self, provider: str, 
                                           current_type: str, 
                                           cpu_score: str, 
                                           memory_score: str) -> Dict:
        """Generar recomendación de rightsizing"""
        if cpu_score in ["muy_bajo", "bajo"] and memory_score in ["muy_bajo", "bajo"]:
            # Recomendar downsizing
            smaller_instance = self._get_smaller_instance(provider, current_type)
            if smaller_instance:
                current_info = self._get_instance_info(provider, current_type)
                new_info = self._get_instance_info(provider, smaller_instance)
                current_cost = current_info['price']
                new_cost = new_info['price']
                savings = (current_cost - new_cost) * 24 * 30  # Ahorro mensual
                
                return {
                    'type': 'rightsizing',
                    'priority': 'high',
                    'action': 'downsize',
                    'current_instance': current_type,
                    'recommended_instance': smaller_instance,
                    'estimated_monthly_savings': round(savings, 2),
                    'description': f"💡 OPTIMIZACIÓN DE RECURSOS: Su instancia {current_type} está subutilizada. Puede reducir costos cambiando a {smaller_instance} sin afectar el rendimiento.",
                    'detailed_analysis': {
                        'current_specs': f"{current_info['vcpu']} vCPU, {current_info['memory']} GB RAM - ${current_cost:.4f}/hora",
                        'recommended_specs': f"{new_info['vcpu']} vCPU, {new_info['memory']} GB RAM - ${new_cost:.4f}/hora",
                        'cost_comparison': f"Actual: ${current_cost * 24 * 30:.2f}/mes → Nuevo: ${new_cost * 24 * 30:.2f}/mes"
                    },
                    'implementation': "🔧 PASOS: 1) Crear snapshot/backup 2) Programar ventana de mantenimiento 3) Cambiar tipo de instancia 4) Verificar funcionamiento",
                    'risk_level': 'low',
                    'pricing_source': self.pricing_sources.get(provider, 'Estimación interna')
                }
        
        elif cpu_score in ["muy_alto"] or memory_score in ["muy_alto"]:
            # Recomendar upsizing
            larger_instance = self._get_larger_instance(provider, current_type)
            if larger_instance:
                return {
                    'type': 'rightsizing',
                    'priority': 'medium',
                    'action': 'upsize',
                    'current_instance': current_type,
                    'recommended_instance': larger_instance,
                    'description': f"Recursos saturados. Cambiar de {current_type} a {larger_instance}",
                    'implementation': "Programar downtime para escalamiento vertical",
                    'risk_level': 'medium'
                }
        
        return None
    
    def _generate_scheduling_recommendation(self, analysis_data: Dict, 
                                          usage_pattern: str) -> Dict:
        """Generar recomendación de programación de instancias"""
        if usage_pattern in ["idle", "variable"]:
            cpu_avg = analysis_data.get('cpu_utilization', {}).get('average', 0)
            potential_savings_pct = 50 if usage_pattern == "idle" else 30
            
            return {
                'type': 'scheduling',
                'priority': 'medium',
                'action': 'automate_start_stop',
                'description': f"⏰ AUTOMATIZACIÓN DE HORARIOS: Su instancia tiene un patrón de uso {usage_pattern} (CPU promedio: {cpu_avg:.1f}%). Puede ahorrar {potential_savings_pct}% automatizando el encendido/apagado.",
                'detailed_analysis': {
                    'usage_pattern': f"Patrón detectado: {usage_pattern}",
                    'efficiency_score': f"Eficiencia actual: {100 - potential_savings_pct}%",
                    'optimal_schedule': "Lunes-Viernes 8:00-18:00, Fines de semana: Apagado"
                },
                'implementation': "🤖 AUTOMATIZACIÓN: 1) AWS: CloudWatch Events + Lambda 2) Azure: Logic Apps + Automation 3) Configurar alertas de verificación 4) Monitorear primer mes",
                'estimated_monthly_savings_percent': potential_savings_pct,
                'schedule_suggestion': {
                    'weekdays': "08:00-18:00 (horario laboral)",
                    'weekends': "Apagado (ahorro 48h)",
                    'holidays': "Apagado (según calendario)"
                },
                'risk_level': 'low',
                'business_impact': 'Minimal - Solo afecta fuera de horario laboral'
            }
        return None
    
    def _generate_reserved_instances_recommendation(self, provider: str,
                                                  instance_type: str,
                                                  usage_pattern: str) -> Dict:
        """Generar recomendación de instancias reservadas"""
        if usage_pattern in ["estable", "intensivo"]:
            savings_1year = 20  # % aproximado de ahorro
            savings_3year = 35
            
            return {
                'type': 'reserved_instances',
                'priority': 'high',
                'action': 'purchase_reserved',
                'description': f"Uso constante detectado. Considerar Reserved Instances para {instance_type}",
                'options': [
                    {
                        'term': '1 year',
                        'payment': 'partial upfront',
                        'estimated_savings': f"{savings_1year}%"
                    },
                    {
                        'term': '3 years', 
                        'payment': 'all upfront',
                        'estimated_savings': f"{savings_3year}%"
                    }
                ],
                'implementation': f"Adquirir Reserved Instance en consola {provider.upper()}",
                'risk_level': 'low'
            }
        return None
    
    def _generate_spot_instances_recommendation(self, provider: str,
                                              instance_type: str,
                                              usage_pattern: str) -> Dict:
        """Generar recomendación de spot instances"""
        if usage_pattern in ["variable"] and provider == "aws":
            return {
                'type': 'spot_instances',
                'priority': 'medium',
                'action': 'migrate_to_spot',
                'description': "Cargas de trabajo tolerantes a interrupciones. Migrar a Spot Instances",
                'estimated_savings': "60-70%",
                'implementation': "Configurar Auto Scaling Group con Spot Instances",
                'considerations': [
                    "Implementar manejo de interrupciones",
                    "Configurar diversificación de tipos de instancia",
                    "Usar mixed instance policy"
                ],
                'risk_level': 'medium'
            }
        return None
    
    def _get_smaller_instance(self, provider: str, current_type: str) -> str:
        """Obtener tipo de instancia más pequeño"""
        if provider == "aws":
            downsize_map = {
                't3.small': 't3.micro',
                't3.medium': 't3.small',
                't3.large': 't3.medium',
                'm5.large': 't3.large',
                'm5.xlarge': 'm5.large',
                'c5.large': 't3.large',
                'c5.xlarge': 'c5.large'
            }
            return downsize_map.get(current_type)
        return None
    
    def _get_larger_instance(self, provider: str, current_type: str) -> str:
        """Obtener tipo de instancia más grande"""
        if provider == "aws":
            upsize_map = {
                't3.micro': 't3.small',
                't3.small': 't3.medium',
                't3.medium': 't3.large',
                't3.large': 'm5.large',
                'm5.large': 'm5.xlarge',
                'c5.large': 'c5.xlarge'
            }
            return upsize_map.get(current_type)
        return None
    
    def _get_instance_cost(self, provider: str, instance_type: str) -> float:
        """Obtener costo por hora de instancia"""
        info = self._get_instance_info(provider, instance_type)
        return info['price']
    
    def _get_instance_info(self, provider: str, instance_type: str) -> Dict:
        """Obtener información completa de instancia"""
        if provider == "aws":
            return self.aws_pricing.get(instance_type, {
                'price': 0.1, 'vcpu': 1, 'memory': 1, 'family': 'Unknown'
            })
        elif provider == "azure":
            return self.azure_pricing.get(instance_type, {
                'price': 0.1, 'vcpu': 1, 'memory': 1, 'family': 'Unknown'
            })
        return {'price': 0.1, 'vcpu': 1, 'memory': 1, 'family': 'Unknown'}
    
    def generate_xlsx_report(self, filepath: str = None) -> str:
        """Generar reporte de optimización en formato XLSX"""
        if not XLSX_AVAILABLE:
            raise ImportError("openpyxl no está instalado. Ejecute: pip install openpyxl")
        
        if not filepath:
            filepath = f"optimon_cost_report_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Reporte de Optimización"
        
        # Estilos
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        center_alignment = Alignment(horizontal="center", vertical="center")
        
        # Encabezados
        headers = ["Instancia", "Tipo Recomendación", "Descripción", "Prioridad", "Ahorro Mensual (USD)", "Implementación", "Riesgo"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment
        
        # Datos
        row = 2
        total_savings = 0
        for rec in self.recommendations_cache:
            ws.cell(row=row, column=1, value=rec.get('instance_id', 'N/A'))
            ws.cell(row=row, column=2, value=rec.get('type', 'N/A').replace('_', ' ').title())
            ws.cell(row=row, column=3, value=rec.get('description', 'N/A'))
            ws.cell(row=row, column=4, value=rec.get('priority', 'N/A').title())
            
            savings = rec.get('estimated_monthly_savings', 0)
            if isinstance(savings, (int, float)):
                ws.cell(row=row, column=5, value=f"${savings:.2f}")
                total_savings += savings
            else:
                ws.cell(row=row, column=5, value="N/A")
            
            ws.cell(row=row, column=6, value=rec.get('implementation', 'N/A'))
            ws.cell(row=row, column=7, value=rec.get('risk_level', 'N/A').title())
            row += 1
        
        # Resumen
        ws.cell(row=row + 1, column=1, value="RESUMEN TOTAL:")
        ws.cell(row=row + 1, column=5, value=f"${total_savings:.2f}")
        
        # Ajustar anchos de columna
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 50
        ws.column_dimensions['D'].width = 12
        ws.column_dimensions['E'].width = 18
        ws.column_dimensions['F'].width = 40
        ws.column_dimensions['G'].width = 12
        
        wb.save(filepath)
        return filepath
    
    def generate_weekly_report(self) -> Dict:
        """Generar reporte semanal de optimización"""
        try:
            report = {
                'report_date': datetime.now().isoformat(),
                'period': 'weekly',
                'summary': {
                    'total_recommendations': len(self.recommendations_cache),
                    'potential_monthly_savings': 0,
                    'high_priority_items': 0
                },
                'recommendations_by_type': {},
                'top_savings_opportunities': [],
                'implementation_roadmap': []
            }
            
            # Procesar recomendaciones
            for rec in self.recommendations_cache:
                rec_type = rec.get('type', 'other')
                if rec_type not in report['recommendations_by_type']:
                    report['recommendations_by_type'][rec_type] = 0
                report['recommendations_by_type'][rec_type] += 1
                
                if rec.get('priority') == 'high':
                    report['summary']['high_priority_items'] += 1
                
                savings = rec.get('estimated_monthly_savings', 0)
                if isinstance(savings, (int, float)):
                    report['summary']['potential_monthly_savings'] += savings
            
            return report
            
        except Exception as e:
            logger.error(f"Error generando reporte semanal: {e}")
            return {}
    
    def get_real_time_metrics(self, instance_id: str, provider: str) -> Dict:
        """Obtener métricas en tiempo real (simulado)"""
        # En implementación real, esto conectaría con CloudWatch/Azure Monitor
        import random
        
        return {
            'cpu_average': round(random.uniform(5, 85), 2),
            'cpu_max': round(random.uniform(50, 95), 2),
            'memory_average': round(random.uniform(10, 80), 2),
            'memory_max': round(random.uniform(40, 90), 2),
            'network_in': random.randint(1000000, 10000000),  # bytes
            'network_out': random.randint(500000, 5000000),   # bytes
            'disk_read_ops': random.randint(100, 1000),
            'disk_write_ops': random.randint(50, 500)
        }

def main():
    """Función principal para testing"""
    engine = CostOptimizationEngine()
    
    # Simular análisis de instancia
    test_metrics = engine.get_real_time_metrics("i-1234567890abcdef0", "aws")
    analysis = engine.analyze_resource_utilization("aws", "i-1234567890abcdef0", test_metrics)
    
    if analysis:
        recommendations = engine.generate_cost_recommendations(analysis, "t3.medium")
        
        print("🎯 ANÁLISIS DE OPTIMIZACIÓN DE COSTOS")
        print("=" * 50)
        print(f"Instancia: {analysis['instance_id']}")
        print(f"CPU Promedio: {analysis['cpu_utilization']['average']}%")
        print(f"Memoria Promedio: {analysis['memory_utilization']['average']}%")
        print(f"Patrón de Uso: {analysis['usage_pattern']}")
        print()
        
        print("💡 RECOMENDACIONES:")
        for i, rec in enumerate(recommendations, 1):
            print(f"{i}. {rec['type'].title()}: {rec['description']}")
            if 'estimated_monthly_savings' in rec:
                print(f"   💰 Ahorro estimado: ${rec['estimated_monthly_savings']}/mes")
            print(f"   🎚️ Prioridad: {rec['priority']}")
            print()

if __name__ == "__main__":
    main()